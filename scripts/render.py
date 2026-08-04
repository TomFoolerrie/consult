#!/usr/bin/env python3
"""
render.py — thin top-level entrypoint the orchestrator invokes to build a Word
deliverable from a component folder (or, for back-compat, a single .md file).

    python3 scripts/render.py <area> -o <out.docx> [--mode working|final]
                              [--slugs a,b] [--track-changes]

For a **folder** (an area dir containing `manifest.json`) this module does the
assembly glue and delegates styling to the bundled CFGI converter:

  1. `doc_model.assemble(folder)` -> structured AssembledDoc (title, subtitle,
     ordered sections each carrying heading / role / slug / number / body).
  2. `doc_model.display_numbers(manifest)` -> the ONE {slug: "L2.seq"} map.
  3. Procedure headings are prefixed with their display number ("1.1 Bank
     Reconciliation"); `[[slug]]` tokens in EVERY section body are resolved to
     display numbers via the same map, so derived tables (Systems "Related
     Procedures", Appendix A "Source Procedure") stay consistent with headings
     after any reorder.
  4. HTML comments (derived markers, scaffold scope notes) and any fenced
     ```consult-meta``` block are stripped so none of them reaches Word.
     All body cleaning is LINE-COUNT-PRESERVING (comments/meta are blanked,
     not removed) so rendered paragraphs map straight back to fragment lines —
     the provenance the M10 review-apply pipeline depends on.
  5. Cover construction branches on input mode: folder -> title/subtitle from
     the manifest and the Document Profile card lifted from the
     `document-profile` static section; single-file -> the converter's legacy
     H1/tagline scan.

M14 — THIS IS ENFORCEMENT POINT 2 FOR THE DOCUMENT PROFILE. Before any other
body cleaning, each procedure body loses the sections the engagement's profile
does not have, the sections its `body_omit` keeps out of the procedure body, and
the callout kinds it does not carry. All three are BLANKED, not deleted, so the
line-for-line provenance mapping survives; with no profile in play they are
no-ops and the output is byte-identical to pre-M14. Fragments are never
rewritten: turning a section back on restores text that was always there.

M16 (moves 3 + 4) — two BODY-VIEW projections applied to each procedure body
right after the profile, both line-count-preserving and both no-ops on a
fragment that uses neither field, so pre-M16 content renders byte-identically:
a callout's `detail:` sub-field is blanked here and rendered only in its
appendix register row (aggregate reads the FRAGMENT, so the detail still reaches
the register when `body_omit` hides its home section), and a step's
`Condition:` tag is hoisted to lead the step body when it was not authored
there.

Modes (M9):
  --mode working (default)  everything visible (gaps, pain points); emits
                            provenance bookmarks + a sidecar review map under
                            {area}/_review/.maps/ keyed by the doc id stamped
                            into the docx core properties.
  --mode final              client-facing: VALIDATION REQUIRED callouts,
                            [[GAP-…]] body flags, and the gap-log appendix are
                            stripped (counts reported, never refused);
                            SCREENSHOT PLACEHOLDER callouts whose captured
                            image exists at _assets/screens/{slug}/{SC-id}.*
                            render as the embedded image with an italic
                            caption below (manual drop-in of a PNG at that
                            path is a first-class alternative to the ingest
                            script). v1.18 client-readiness grooming: the
                            screenshot-index appendix never ships (evidence
                            lives inline; pending count goes to the
                            scorecard), register lead-in prose is stripped,
                            zero-row registers are skipped whole, lexicon
                            terms (_client/lexicon.yaml) normalize spelling
                            variants, and a READINESS scorecard reports
                            placeholders / blank required register cells /
                            doubled spaces / doubled punctuation — counts
                            reported, never refused; --strict exits 1 when
                            the scorecard is dirty. Never writes the
                            .render.json signal (M21):
                            final mode is a terminal export produced FROM an
                            accepted state, not a transition into a new one —
                            same reasoning as --slugs below.
  --slugs a,b               subset render: only those procedure sections (no
                            cover/front/back matter) — the per-owner kit docs.
                            Display numbers + callout display IDs still come
                            from the FULL manifest so kit docs match the full
                            draft. Never writes the .render.json signal.
  --track-changes           kit docs open with tracked changes already on.

Folder renders (both modes, v1.18): a blank fill-by-hand Document Control
table opens the front matter; the TOC field collects H1-H2 and settings.xml
carries <w:updateFields/> so Word populates it on open; every Reference &
Appendices section after the first opens a fresh page.

The assembled Markdown/structure is then handed to
`cfgi_markdown_to_word.convert_assembled` (folder) or `.convert` (single file).
`scripts/doc_model.py` is owned by M2 and imported here, never created.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import sys

import console_compat  # noqa: F401  (stdout errors='replace' on narrow consoles)
import uuid
from pathlib import Path

# Make both scripts/ (doc_model) and the skill's scripts/ (the converter)
# importable regardless of the caller's cwd.
_SCRIPTS_DIR = Path(__file__).resolve().parent
_REPO_ROOT = _SCRIPTS_DIR.parent
_SKILL_SCRIPTS = _REPO_ROOT / "skills" / "consult-docx-builder" / "scripts"
for _p in (_SCRIPTS_DIR, _SKILL_SCRIPTS):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import doc_model  # noqa: E402  (M2-owned shared spine)
import client_config  # noqa: E402  (M13 resolver / M14 document profile)
import callouts  # noqa: E402  (shared callout grammar — M16 note/detail fields)
import cfgi_markdown_to_word as cfgi  # noqa: E402  (bundled converter)

from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

# Static section(s) lifted onto the cover as the Document Profile card. Reuse
# the converter's set so the two stay in lock-step.
_PROFILE_HEADINGS = set(cfgi.COVER_SECTIONS)  # {"document profile"}

# Where captured screenshots live (per procedure, keyed by LOCAL SC id).
SCREENS_DIR = "_assets/screens"
_IMG_EXTS = (".png", ".jpg", ".jpeg", ".gif")


# --------------------------------------------------------------------------- #
# doc_model access shims (tolerant to dict- or dataclass-shaped returns and to
# resolve_tokens' exact signature, which M2 finalizes).
# --------------------------------------------------------------------------- #
def _attr(obj, name, default=None):
    if isinstance(obj, dict):
        return obj.get(name, default)
    return getattr(obj, name, default)


def _sections(assembled):
    secs = _attr(assembled, "sections")
    if secs is None:
        secs = _attr(assembled, "components")
    return secs or []


def _resolve_tokens(text: str, numbers) -> str:
    """Resolve `[[slug]]` cross-references to display numbers via doc_model."""
    try:
        return doc_model.resolve_tokens(text, numbers, "number")
    except TypeError:
        # M2 may expose resolve_tokens without a mode arg.
        return doc_model.resolve_tokens(text, numbers)


# --------------------------------------------------------------------------- #
# Body cleaning (things that must never reach Word) — LINE-COUNT-PRESERVING
# --------------------------------------------------------------------------- #
def _strip_html_comments(text: str) -> str:
    """Blank ALL `<!-- ... -->` comments (derived markers, scaffold scope
    notes, stray sentinels) — authoring metadata, never content. Multi-line
    comments are replaced by the same number of newlines so provenance line
    mapping survives."""
    return re.sub(r"<!--.*?-->",
                  lambda m: "\n" * m.group(0).count("\n"),
                  text, flags=re.DOTALL)


def _strip_consult_meta(text: str) -> str:
    """Blank any fenced block whose info-string is `consult-meta` (```/~~~)."""
    lines = text.split("\n")
    out = []
    i = 0
    open_re = re.compile(r"^\s*(`{3,}|~{3,})\s*consult-meta\s*$", re.IGNORECASE)
    while i < len(lines):
        m = open_re.match(lines[i])
        if m:
            fence_char = m.group(1)[0]
            close_re = re.compile(r"^\s*" + re.escape(fence_char) + r"{3,}\s*$")
            out.append("")
            i += 1
            while i < len(lines) and not close_re.match(lines[i]):
                out.append("")
                i += 1
            if i < len(lines):
                out.append("")   # the closing fence
                i += 1
            continue
        out.append(lines[i])
        i += 1
    return "\n".join(out)


_CALLOUT_ID_RE = re.compile(r"\b(?:CTRL|GAP|PP|IO|SC)-[A-Z0-9]+(?:-[A-Z0-9]+)*\b")

#: Citation scrub (final mode only). SRC-## ids are the drafters' provenance
#: citations and GAP-## ids point at callouts/log rows that final mode strips —
#: neither belongs in a client-facing export. Only two UNAMBIGUOUS shapes are
#: removed mechanically: a parenthetical containing NOTHING but citation ids
#: ("(SRC-002, SRC-005)", "(see GAP-011)") and a pure-citation sentence
#: ("See GAP-012 and GAP-014."). An id woven into sentence meaning ("; see
#: GAP-011, which is unresolved") is prose — it is never auto-edited; it falls
#: through to the dangling-reference warning for a human to reword.
#: Fragments hard-wrap prose at ~80 columns, so a citation may span a line
#: break; _WS allows AT MOST ONE newline so a match can rejoin a wrapped
#: paragraph but can never swallow a blank line (a paragraph boundary).
_WS = r"[ \t]*(?:\n[ \t]*)?"
_CITE_ID = r"(?:SRC|GAP)-\d+"
_CITE_PAREN_RE = re.compile(
    rf"{_WS}\({_WS}(?:see{_WS})?{_CITE_ID}"
    rf"(?:{_WS}(?:[,;]|\band\b){_WS}{_CITE_ID})*{_WS}\)",
    re.IGNORECASE)
_CITE_SENTENCE_RE = re.compile(
    rf"{_WS}\bSee{_WS}{_CITE_ID}"
    rf"(?:{_WS}(?:[,;]|\band\b){_WS}{_CITE_ID})*{_WS}\.")
#: What the scrub could not remove — the final-mode dangling-reference detector
#: (formerly GAP-only; SRC ids dangle identically once sources.yaml is not
#: shipped with the deliverable).
_GAP_ID_RE = re.compile(rf"\b{_CITE_ID}\b")


def _scrub_citations(text: str, stats: dict) -> str:
    text, n1 = _CITE_PAREN_RE.subn("", text)
    text, n2 = _CITE_SENTENCE_RE.subn("", text)
    # Sentence skeletons left when the [[GAP-…]] tag strip deleted the whole
    # object of a "see" clause: "owner. See ." / "owner. See and ." (the tag
    # opened its own sentence) and "unconfirmed — see." / "— see ." (a
    # trailing dash clause). Then seam repair: a space stranded before
    # punctuation and doubled interior spaces — both anchored on a following
    # non-space so a trailing double space (markdown hard break) survives.
    text = re.sub(rf"(?:(?<=[.!?])|(?<=\n))[ \t]*\bSee(?:[ \t]+and\b)*"
                  rf"{_WS}[,;]?[ \t]*\.", "", text)
    text = re.sub(rf"[ \t]*[—–-][ \t]*\bsee(?:[ \t]+and\b)*{_WS}[,;]?[ \t]*"
                  rf"(?=\.)", "", text, flags=re.IGNORECASE)
    if n1 or n2:
        text = re.sub(r" +([.,;:!?])", r"\1", text)
        text = re.sub(r"(?<=\S)  +(?=\S)", " ", text)
    stats["citations_scrubbed"] += n1 + n2
    return text


def _rewrite_callout_ids(text: str, submap: dict) -> str:
    """Rewrite one procedure's local callout IDs to their global display IDs.

    submap is {local-id: display-id} for THIS procedure only (context makes the
    local IDs unambiguous). Single-pass callback substitution: every match is
    resolved against the original text, so old/new ID overlaps cannot chain.
    Unknown IDs (e.g. a reference reconcile will flag) pass through untouched.
    """
    return _CALLOUT_ID_RE.sub(lambda m: submap.get(m.group(0), m.group(0)), text)


def _flag_gap_tags(text: str) -> str:
    """Render body gap tags `[[GAP-NN — label]]` as a visible bold flag.

    doc_model.resolve_tokens deliberately skips them (they are not procedure
    cross-refs), so without this they reach Word as raw wiki brackets.
    """
    return re.sub(r"\[\[\s*(GAP-\d+[^\]]*?)\s*\]\]", r"**[\1]**", text)


def _strip_own_heading(text: str) -> str:
    """Blank the fragment's own leading `## Heading` line (line preserved).

    Each component file opens with its own H2 (the single-H1/flat-H2 contract),
    but the assembled doc emits the manifest heading — the numbering/TOC
    authority — so the body's copy would render every section heading twice.
    Only a leading H2 is blanked; anything after real content is left alone.
    """
    lines = text.split("\n")
    for idx, ln in enumerate(lines):
        if ln.strip():
            if ln.startswith("## "):
                lines[idx] = ""
            break
    return "\n".join(lines)


def _clean_body(text: str, numbers) -> str:
    text = _strip_own_heading(text)
    text = _strip_consult_meta(text)
    text = _strip_html_comments(text)
    text = _resolve_tokens(text, numbers)
    text = _flag_gap_tags(text)
    return text


# --------------------------------------------------------------------------- #
# M14 profile enforcement (enforcement point 2) — LINE-COUNT-PRESERVING
#
# Sections the profile does not have, sections `body_omit` keeps out of the
# procedure body, and callout kinds the profile does not carry are BLANKED, not
# deleted: exactly the discipline the comment/consult-meta strippers above
# follow. Blank lines produce no Word output, so the reader sees the section
# gone, while assembled body line k still maps to fragment line k — the
# provenance the M10 review-apply pipeline depends on. (The final-mode
# transforms below may drop lines because final mode has no review round; a
# `body_omit` render is an ordinary WORKING render and does.)
#
# Both are no-ops that return the text UNCHANGED when the profile excludes
# nothing, so an engagement with no `profile.yaml` renders byte-identically.
# --------------------------------------------------------------------------- #
_META_FENCE_RE = re.compile(r"^\s*(`{3,}|~{3,})\s*consult-meta\s*$", re.IGNORECASE)

# A callout label line, loose enough to catch every kind in the LABEL map:
# `> **<LABEL> — <ID>:** text`. The grammar itself is callouts.py's; this only
# needs the label to decide whether the block stays.
_CALLOUT_LABEL_RE = re.compile(
    r"^\s*>\s*\*\*\s*(?P<label>[A-Z][A-Z ]*?)\s*[-–—]\s*[A-Za-z]+-[A-Za-z0-9\-]+\s*:"
)


def _blank_sections(text: str, slugs) -> str:
    """Blank the `###` blocks of every section slug in `slugs` (M23).

    Keyed on the section SLUG, resolved from the heading by
    `doc_model.section_slug` — which accepts the migrated form
    (`### Key Controls`) and the not-yet-migrated one (`### F. Key Controls`)
    alike, so an unmigrated area keeps rendering exactly as it does today.

    The `consult-meta` end matter belongs to no section, so it survives even
    when the section it trails is hidden — otherwise hiding `issues` would strip
    a procedure's system/role bindings out of the assembled body.
    """
    if not slugs:
        return text
    out: list[str] = []
    dropping = False
    for ln in text.split("\n"):
        if _META_FENCE_RE.match(ln):
            dropping = False
        else:
            slug = doc_model.section_of_heading(ln)
            if slug is not None:
                dropping = slug in slugs
        out.append("" if dropping else ln)
    return "\n".join(out)


def _letter_sections(text: str, letters) -> str:
    """Stamp the display LETTER onto every section heading (M23).

    The identity/display split for sections: the fragment carries
    `### Scope`, the DOCUMENT carries `### A. Scope`, and
    `letters` (the profile's `sections:` order) decides which letter. Exactly
    what `_heading_for` does with a procedure's display number, and exactly why
    reordering a profile re-letters a render with zero fragment edits.

    LINE-COUNT-PRESERVING (one heading line rewritten in place) and it runs
    BEFORE the body is emitted, so the M10 provenance map — whose entry hashes
    are taken from RENDERED paragraph text — sees the letters, not the titles.

    A heading that already carries a letter has it REPLACED, so an unmigrated
    fragment renders with the letter the profile says, never the stale one baked
    into its heading.

    The TITLE comes from the registry whenever the heading was resolved BY title
    (which is what makes a registry rename show up in the document without any
    fragment edit), and is left exactly as authored when only the LETTER
    resolved it — a heading the registry does not know is local wording this
    transform has no business rewriting.
    """
    out = []
    seen: set[str] = set()
    for ln in text.split("\n"):
        slug = doc_model.section_of_heading(ln)
        if slug is None or slug not in letters:
            out.append(ln)
            continue
        written = doc_model.section_heading_title(ln)
        title = (doc_model.section_title(slug)
                 if doc_model.section_of_title(written) == slug else written)
        # M16 move 1's C+D merge, pre-content-wave: a fragment carrying BOTH
        # `### Pre-Requisites` and `### Inputs` resolves both to
        # `before-you-start`. Only the FIRST takes the registry title; a repeat
        # keeps its authored one under the same letter, so nothing is lost and
        # the un-merged pair is visible rather than printed twice identically
        # (see doc_model.duplicate_sections — tolerate + report, never fail).
        if slug in seen:
            title = written
        seen.add(slug)
        out.append(f"### {letters[slug]}. {title}")
    return "\n".join(out)


def _blank_callouts(text: str, labels) -> str:
    """Blank every callout block whose LABEL is in `labels`.

    A callout is its label line plus the contiguous `>` blockquote under it —
    the same block shape `_final_transform` consumes for VALIDATION REQUIRED,
    which is that stripper generalized to any kind the profile drops.
    """
    if not labels:
        return text
    lines = text.split("\n")
    out: list[str] = []
    i = 0
    while i < len(lines):
        m = _CALLOUT_LABEL_RE.match(lines[i])
        label = re.sub(r"\s+", " ", m.group("label")).strip() if m else None
        if label in labels:
            out.append("")
            i += 1
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                if _CALLOUT_LABEL_RE.match(lines[i]):
                    break      # a new callout starts its own block
                out.append("")
                i += 1
            continue
        out.append(lines[i])
        i += 1
    return "\n".join(out)


def _apply_profile(text: str, profile) -> str:
    """Strip everything the document profile excludes from one procedure body,
    then stamp the section letters the profile's order assigns (M23).

    Hiding first, lettering second: a section's letter is its position in
    `sections:`, which `body_omit` does not change, so hiding one never shifts
    another's letter.
    """
    text = _blank_sections(text, profile.hidden_sections())
    text = _blank_callouts(text, profile.dropped_callouts())
    return _letter_sections(text, profile.letters())


# --------------------------------------------------------------------------- #
# M16 move 3 — callout `detail:` is appendix-only  (LINE-COUNT-PRESERVING)
#
# One source of truth, two views: the fragment holds the whole callout, the
# STEP shows `note:` (one or two sentences) and the appendix register row shows
# `detail:` (the full account). This is the projection half — blanking the
# `detail:` bullet and its wrapped continuation lines out of the procedure body.
# aggregate.py is the other half and reads the FRAGMENT, so a detail still
# reaches its register when the profile has body_omit'ed its home section: the
# two features compose without knowing about each other.
#
# A callout with no `detail:` is untouched, so a fragment using none of this
# grammar renders byte-identically to pre-M16.
# --------------------------------------------------------------------------- #
def _hide_callout_details(text: str) -> str:
    """Blank the `> - **Detail:** …` bullet of every callout (body view only)."""
    lines = text.split("\n")
    out = list(lines)
    in_callout = False
    i = 0
    while i < len(lines):
        ln = lines[i]
        if _CALLOUT_LABEL_RE.match(ln):
            in_callout = True
            i += 1
            continue
        if not ln.lstrip().startswith(">"):
            in_callout = False
            i += 1
            continue
        if in_callout and callouts.is_callout_field(ln, callouts.DETAIL_FIELD):
            # Blank to a bare `>` — NOT to an empty line. The line count is
            # preserved either way (provenance), but a fully blank line ends
            # the blockquote, so the converter would split one callout into
            # two boxes around the hidden detail (label + note above, the
            # remaining sub-fields below, a page gap between). A bare quote
            # marker keeps the block contiguous and renders nothing.
            out[i] = ">"
            i += 1
            # A wrapped value continues over plain `>` lines until the next
            # sub-field bullet, the next callout, or the end of the block —
            # the same continuation rule aggregate's parser applies.
            while i < len(lines):
                nxt = lines[i]
                if (not nxt.lstrip().startswith(">")
                        or _CALLOUT_LABEL_RE.match(nxt)
                        or callouts.callout_field(nxt)):
                    break
                out[i] = ">"
                i += 1
            continue
        i += 1
    return "\n".join(out)


# --------------------------------------------------------------------------- #
# M16 move 4 — a step's `Condition:` LEADS its body  (LINE-COUNT-PRESERVING)
#
# The contract asks drafters to author the `Condition:` tag as the first line of
# the step body, directly under the heading, because that is where it must be
# read: "Step 5" that reads as "next, do this" is the defect. This transform is
# the guarantee for a fragment authored the other way round (the tag sitting
# with its sibling inline tags below the prose) — it ROTATES the tag up to the
# head of the step body, which keeps the line COUNT identical.
#
# It fires only when the tag is not already leading, so a contract-compliant
# fragment is passed through untouched and its M10 line-for-line provenance is
# exact. When it does fire, provenance shifts by one line inside that one step —
# the price of rendering a mis-ordered fragment correctly, paid nowhere else.
# --------------------------------------------------------------------------- #
_STEP_HEADING_RE = re.compile(r"^\s{0,3}#{4,6}\s+\S")
_ANY_HEADING_RE = re.compile(r"^\s{0,3}#{1,6}\s+\S")
_CONDITION_RE = re.compile(
    r"^\s*[-*+]\s*\*\*\s*" + re.escape(client_config.CONDITION_TAG)
    + r"\s*:\s*\*\*", re.IGNORECASE)


def _lead_conditions(text: str) -> str:
    """Hoist each step's `- **Condition:** …` tag to the head of its body."""
    lines = text.split("\n")
    n = len(lines)
    starts = [i for i, ln in enumerate(lines) if _STEP_HEADING_RE.match(ln)]
    for h in starts:
        # The step body runs to the next heading of any level.
        end = n
        for j in range(h + 1, n):
            if _ANY_HEADING_RE.match(lines[j]):
                end = j
                break
        cond = None
        for j in range(h + 1, end):
            if lines[j].lstrip().startswith(">"):
                continue          # a callout's own bullets are not step tags
            if _CONDITION_RE.match(lines[j]):
                cond = j
                break
        if cond is None:
            continue
        first = next((j for j in range(h + 1, end) if lines[j].strip()), None)
        if first is None or first == cond:
            continue              # already leading — nothing to do
        if lines[h + 1].strip():
            # No blank line under the heading to rotate into: hoisting here
            # would let the prose be absorbed into the tag's own bullet
            # paragraph. Leave it in place rather than corrupt the step.
            continue
        rotated = [lines[cond]] + lines[h + 1:cond]
        lines[h + 1:cond + 1] = rotated
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# RACI display ordering — like letters and display numbers, row order is a
# render-time concern: the derived file's authored order is never rewritten.
# --------------------------------------------------------------------------- #
_RACI_ROW_SLUG_RE = re.compile(r"^\|\s*\[\[\s*([a-z0-9][a-z0-9-]*)\s*\]\]")
_TABLE_SEP_RE = re.compile(r"^\s*\|?[\s|:\-]+\|?\s*$")


def _order_raci_rows(text: str, numbers) -> str:
    """Sort the RACI matrix's activity rows into procedure display-number
    order (1.1, 1.2, … 2.1). Fires only when EVERY row of the table leads
    with a resolvable `[[slug]]` token — anything else (a hand-added row, an
    unknown slug) leaves the authored order untouched rather than guessing."""
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        if not (ln.lstrip().startswith("|") and i + 1 < len(lines)
                and "-" in lines[i + 1] and _TABLE_SEP_RE.match(lines[i + 1])):
            continue
        j = i + 2
        rows = []
        while j < len(lines) and lines[j].lstrip().startswith("|"):
            rows.append(lines[j])
            j += 1
        keys = []
        for r in rows:
            m = _RACI_ROW_SLUG_RE.match(r.lstrip())
            num = numbers.get(m.group(1)) if m else None
            if not num:
                return text
            keys.append(tuple(int(x) for x in num.split(".")))
        ordered = [r for _, r in sorted(zip(keys, rows), key=lambda t: t[0])]
        return "\n".join(lines[:i + 2] + ordered + lines[j:])
    return text


# --------------------------------------------------------------------------- #
# Final-mode transforms (line count NOT preserved — final mode has no
# provenance; there is no review round against a final deliverable)
# --------------------------------------------------------------------------- #
_VR_LINE = re.compile(r"^\s*>\s*\*\*\s*VALIDATION REQUIRED\b")
_SC_LINE = re.compile(
    r"^\s*>\s*\*\*\s*SCREENSHOT PLACEHOLDER\s*[-–—]\s*"
    r"(?P<id>SC-[A-Z0-9]+(?:-[A-Z0-9]+)*)\s*:\*\*\s*(?P<text>.*)$"
)


def _find_screen(folder: Path, slug: str, local_id: str) -> Path | None:
    base = folder / SCREENS_DIR / (slug or "")
    for ext in _IMG_EXTS:
        p = base / f"{local_id}{ext}"
        if p.is_file():
            return p
    return None


def _final_transform(text: str, slug: str, folder: Path,
                     disp_to_local: dict, stats: dict) -> str:
    """Strip open-gap scaffolding; swap captured screenshots in for their
    placeholder boxes. Runs AFTER the display-ID rewrite, so SC labels carry
    display ids here; `disp_to_local` maps them back to the asset filename."""
    out: list[str] = []
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        if _VR_LINE.match(ln):
            stats["gaps_stripped"] += 1
            i += 1
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                i += 1
            continue
        m = _SC_LINE.match(ln)
        if m:
            block_start = i
            disp_id, desc = m.group("id"), m.group("text").strip()
            i += 1
            extra: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                t = lines[i].lstrip().lstrip(">").strip()
                if t and not t.startswith("- **"):
                    extra.append(t)
                i += 1
            local = disp_to_local.get(disp_id, disp_id)
            img = _find_screen(folder, slug, local)
            if img:
                stats["screens_embedded"] += 1
                caption = " ".join([f"{disp_id}:", desc] + extra).strip()
                out.append(f"![{caption}]({img})")
            else:
                stats["screens_placeholder"] += 1
                out.extend(lines[block_start:i])
            continue
        out.append(ln)
        i += 1
    # Inline body gap tags vanish in final mode. Whole-text, because fragments
    # hard-wrap prose: the tag (and the whitespace before it) may span a line
    # break; at most one newline is consumed so a blank line (paragraph
    # boundary) never is.
    joined, n = re.subn(r"[ \t]*(?:\n[ \t]*)?\[\[[ \t]*GAP-[^\]]*\]\]", "",
                        "\n".join(out))
    stats["gap_tags_stripped"] += n
    return joined


# --------------------------------------------------------------------------- #
# Final-mode register grooming + readiness scorecard (v1.18)
# --------------------------------------------------------------------------- #
_ITALIC_LINE_RE = re.compile(r"^_[^_].*_$")


def _strip_register_leadin(text: str) -> str:
    """Drop a derived section's leading italic context sentence (final mode
    only). The lead-ins orient a REVIEWER ("canonical roles, with the
    procedures each appears in"); between the heading and the table the client
    wants the table."""
    lines = text.split("\n")
    for idx, ln in enumerate(lines):
        if not ln.strip():
            continue
        if _ITALIC_LINE_RE.match(ln.strip()):
            del lines[idx]
        break
    return "\n".join(lines)


def _table_data_rows(text: str) -> int:
    """Count table DATA rows (not headers, not separators) in a body."""
    lines = text.split("\n")
    n = 0
    for i, ln in enumerate(lines):
        if not ln.lstrip().startswith("|") or _TABLE_SEP_RE.match(ln):
            continue
        if i + 1 < len(lines) and "-" in lines[i + 1] \
                and _TABLE_SEP_RE.match(lines[i + 1]):
            continue          # header row (a separator sits directly under it)
        n += 1
    return n


#: Register columns a client deliverable must not ship blank, by derived kind:
#: {cell index: column name} (index into the row's cells, 0 = first column).
_REQUIRED_COLS = {
    "role-dictionary": {1: "Reports To"},
    "systems": {1: "Role in Process"},
    "appendix-a": {1: "Impact", 2: "Severity"},
}

_PLACEHOLDER_RE = re.compile(
    r"\bTBD\b|Pending user input|Pending synthesis", re.IGNORECASE)
_DOUBLE_SPACE_RE = re.compile(r"(?<=\S)  +(?=\S)")
_DOUBLE_PUNCT_RE = re.compile(r"(?<!\.)\.\.(?!\.)|,,|;;")


def _readiness_scan(body: str, where: str, derived_kind, readiness: dict) -> None:
    """Final-mode readiness checks, run on what the client would read.

    Counts reported, never refused — the render always completes; `--strict`
    lets a caller turn a dirty scorecard into a nonzero exit."""
    blanks = _REQUIRED_COLS.get(derived_kind or "")
    lines = body.split("\n")
    for i, ln in enumerate(lines):
        s = ln.strip()
        m = _PLACEHOLDER_RE.search(ln)
        if m:
            readiness["placeholders"].append(f"{where}: {m.group(0)!r}")
        if s.startswith("|"):
            if blanks and not _TABLE_SEP_RE.match(ln) and not (
                    i + 1 < len(lines) and "-" in lines[i + 1]
                    and _TABLE_SEP_RE.match(lines[i + 1])):
                cells = [c.strip() for c in s.strip("|").split("|")]
                for idx, col in blanks.items():
                    if idx < len(cells) and cells[idx] in ("", "—", "-"):
                        row = cells[0][:48] if cells else "?"
                        readiness["blank_cells"].append(
                            f"{where}: blank {col!r} ({row}…)")
            continue
        if s.startswith("#"):
            continue
        if _DOUBLE_SPACE_RE.search(ln.rstrip()):
            readiness["double_spaces"].append(f"{where}: {s[:60]!r}")
        if _DOUBLE_PUNCT_RE.search(ln):
            readiness["double_punct"].append(f"{where}: {s[:60]!r}")


def _lexicon_normalizers(folder: Path):
    """Compiled (pattern, canonical) pairs from `_client/lexicon.yaml`."""
    pairs = []
    for term in client_config.lexicon(folder):
        pairs.append((re.compile(rf"(?<!\w){re.escape(term)}(?!\w)",
                                 re.IGNORECASE), term))
    return pairs


def _apply_lexicon(text: str, normalizers, stats: dict) -> str:
    """Normalize case-variants of lexicon terms to their canonical spelling.
    Idempotent: an already-canonical match is left alone (and not counted)."""
    for rx, canon in normalizers:
        def _sub(m):
            if m.group(0) == canon:
                return m.group(0)
            stats["lexicon_normalized"] += 1
            return canon
        text = rx.sub(_sub, text)
    return text


# --------------------------------------------------------------------------- #
# Provenance (M9 stamps it; M10 consumes it)
# --------------------------------------------------------------------------- #
def _add_bookmark(p, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    pPr = p._p.find(qn("w:pPr"))
    p._p.insert(1 if pPr is not None else 0, start)
    p._p.append(end)


class Prov:
    """Provenance collector handed to the converter's render_body.

    `origins` is a list parallel to the assembled body lines; each entry is
    None (heading/derived glue) or {file, slug, l2, line} where `line` is the
    0-based line index in the SOURCE fragment (cleaning preserves counts, so
    assembled body line k of a section ↔ fragment file line k).
    """

    def __init__(self):
        self.origins: list = []
        self.entries: dict[str, dict] = {}
        self._n = 0

    def mark(self, paras, i0: int, i1: int, kind: str) -> None:
        o = self.origins[i0] if 0 <= i0 < len(self.origins) else None
        oe = self.origins[i1 - 1] if 0 <= i1 - 1 < len(self.origins) else None
        if not o or not oe or oe["file"] != o["file"]:
            return
        for sub, p in enumerate(paras):
            name = f"cw_{self._n:05d}"
            _add_bookmark(p, name, 9000 + self._n)
            self._n += 1
            self.entries[name] = {
                "file": o["file"],
                "slug": o.get("slug"),
                "l2": o.get("l2"),
                "lines": [o["line"], oe["line"]],
                "kind": kind,
                "sub": sub,
                "hash": hashlib.sha1(p.text.encode("utf-8")).hexdigest(),
            }


def _write_review_map(folder: Path, out: Path, doc_id: str, prov: Prov,
                      basis: dict) -> Path:
    maps_dir = folder / "_review" / ".maps"
    maps_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": "consult-review-map/v1",
        "doc_id": doc_id,
        "area": folder.name,
        "docx": out.name,
        "basis": basis,
        "entries": prov.entries,
    }
    map_path = maps_dir / f"{doc_id}.json"
    map_path.write_text(json.dumps(payload, indent=1, ensure_ascii=False) + "\n",
                        encoding="utf-8")
    return map_path


# --------------------------------------------------------------------------- #
# Folder assembly glue
# --------------------------------------------------------------------------- #
def _heading_for(section, numbers) -> str:
    heading = (_attr(section, "heading") or "").strip()
    if _attr(section, "role") == "procedure":
        num = _attr(section, "number")
        if not num:
            slug = _attr(section, "slug")
            num = numbers.get(slug) if slug else None
        if num:
            return f"{num} {heading}"
    return heading


def render_folder(folder: Path, out: Path, *,
                  landscape: bool = False, do_cover: bool = True,
                  mode: str = "working", slugs: list[str] | None = None,
                  track_changes: bool = False, emit_signal: bool = True) -> dict:
    """Render an area folder. Returns a stats dict (counts + doc_id/map)."""
    folder = Path(folder)
    # M14 enforcement point 2. Resolved before any body is touched so a
    # malformed profile fails the render rather than shipping a wrong shape.
    profile = client_config.profile(folder)
    manifest = doc_model.load_manifest(folder)
    doc_model.validate_manifest(manifest)
    # The profile validator's cross-field rule refuses a `body_omit:` section
    # whose register is not in `derived:` — but it can only vouch for the
    # PROFILE. Registers are manifest components, written at the confirm gate
    # (enforcement point 1); a profile that acquired the omission after that
    # gate ran would blank the section from every body below while the
    # manifest lists no register to catch the callouts, and they would
    # silently leave the document. Same table, same rule, re-checked here
    # against what this render actually builds from.
    manifest_registers = {c.get("derived_kind")
                          for c in manifest.get("components", [])}
    for section, (register, what) in client_config.BODY_OMIT_REGISTERS.items():
        if section in profile.body_omit and register not in manifest_registers:
            raise SystemExit(
                f"error: the profile's `body_omit:` hides the `{section}` "
                f"section from the procedure body, but this manifest lists no "
                f"`{register}` component to catch the {what} — rendering "
                f"would drop them from the document. The manifest predates "
                f"the profile change: run `python3 scripts/scaffold.py "
                f"--sync-profile --area {folder}` to add the register, "
                f"re-aggregate, then re-render."
            )
    # Callout kinds with no HOME at all — the profile deliberately drops the
    # kind's home section from `sections:` (or the kind itself from
    # `callouts:`) with no register in the manifest. Sanctioned, unlike the
    # body_omit trap refused above: it declares "this document carries no
    # controls / pain points". But a prose mention of such an id ("...the
    # assurance provided by CTRL-31") then dangles for the reader exactly
    # like a final-mode gap reference — and in EVERY mode, because the hole
    # is profile-shaped, not export-shaped. Driven entirely by the shared
    # registries (label -> home section -> register), so a new kind or
    # register is covered by its table rows.
    homeless_prefixes = sorted(
        callouts.LABEL_TO_PREFIX[label]
        for label, home in callouts.LABEL_TO_HOME_SECTION.items()
        if home in client_config.BODY_OMIT_REGISTERS
        and client_config.BODY_OMIT_REGISTERS[home][0] not in manifest_registers
        and (home in profile.dropped_sections()
             or label in profile.dropped_callouts()))
    homeless_re = (re.compile(r"\b(?:%s)-\d+\b" % "|".join(homeless_prefixes))
                   if homeless_prefixes else None)
    numbers = doc_model.display_numbers(manifest)
    # Token resolution map: [[slug]] -> "1.1 Vendor Onboarding". Bare numbers
    # are correct for section headings, but in prose and derived tables (where
    # procedure-local IDs repeat across rows) the number alone is cryptic —
    # number + title makes every cross-reference self-explanatory.
    labels = dict(numbers)
    l2_of: dict[str, str] = {}
    for comp in manifest.get("components", []):
        slug = comp.get("slug")
        if slug and slug in labels and comp.get("heading"):
            labels[slug] = f"{labels[slug]} {comp['heading']}"
        if slug:
            l2_of[slug] = comp.get("l2", "")
    # M26 cross-area tokens: [[area/slug]] -> "Heading (Area Title)". Under a
    # components/ root every sibling procedure resolves; elsewhere the map is
    # empty and a cross token fails loud (KeyError) rather than reaching Word
    # as raw brackets — reconcile reports the same defect first.
    labels.update(doc_model.cross_labels(folder))
    assembled = doc_model.assemble(folder)

    # Global callout display IDs (drafters number locally from 01; the global
    # sequence is a render-time transform — see doc_model.callout_display_ids).
    # ALWAYS computed over the full folder, so subset (kit) docs and derived
    # views agree with the full draft.
    id_map = doc_model.callout_display_ids(folder)
    ids_by_slug: dict = {}
    disp_to_local_by_slug: dict = {}
    for (slug, local), disp in id_map.items():
        ids_by_slug.setdefault(slug, {})[local] = disp
        disp_to_local_by_slug.setdefault(slug, {})[disp] = local

    title = _attr(assembled, "title") or ""
    subtitle = _attr(assembled, "subtitle") or ""
    stats = {"mode": mode, "gaps_stripped": 0, "gap_tags_stripped": 0,
             "citations_scrubbed": 0,
             "dangling_gap_refs": {}, "dangling_gap_ref_count": 0,
             "dangling_refs": {}, "dangling_ref_count": 0,
             "screens_embedded": 0, "screens_placeholder": 0,
             "lexicon_normalized": 0, "empty_registers_skipped": [],
             "readiness": {"placeholders": [], "blank_cells": [],
                           "double_spaces": [], "double_punct": [],
                           "cover_fields": []},
             "profile": profile.report_line()}

    lexicon_normalizers = _lexicon_normalizers(folder) if mode == "final" else []
    subset = slugs is not None
    if subset:
        unknown = [s for s in slugs if s not in numbers]
        if unknown:
            raise SystemExit(f"error: unknown procedure slug(s): {', '.join(unknown)}")
        do_cover = False   # kit docs are lean: content starts at the heading

    provenance = (mode == "working")
    prov = Prov() if provenance else None

    profile_md = ""
    body_lines: list[str] = []
    origins: list = []

    def emit(line: str, origin=None):
        body_lines.append(line)
        origins.append(origin)

    # L2 chapter dividers — display glue, exactly like the numbers themselves.
    # The body gives sub-process boundaries no signal (1.3 is followed by 2.1
    # with nothing between) while the appendix registers already group their
    # rows by L2: the body claims a hierarchy it never shows. Before the first
    # procedure of each bucket this emits `# {ordinal}. {Title}` — ordinal from
    # the same l2_order index display_numbers uses, title by the same titleize
    # rule aggregate's l2_titles applies — and one `# Reference & Appendices`
    # before the back matter, so the TOC does not nest the registers under the
    # last sub-process. H1 is otherwise unused in folder docs; the converter
    # gives it a page break. Kit (subset) docs carry no dividers: a chapter
    # head over a single excerpted procedure is noise.
    l2_ordinal = {l2: i for i, l2 in
                  enumerate(manifest.get("l2_order", []) or [], start=1)}
    divider_state = {"l2": None, "front": False, "in_procedures": False,
                     "backmatter": False}
    # Reference & Appendices sections each open a fresh page — EXCEPT the
    # first, which stays under the chapter heading (breaking it too would
    # strand the H1 alone on a page). Heading texts, matched by the converter.
    break_headings: set[str] = set()

    # Document Control front matter: a blank fill-by-hand table under its own
    # chapter heading, straight after the cover + TOC. Emitted glue, never a
    # fragment — there is nothing to draft, only version history to record.
    if do_cover and not subset:
        emit("# Document Control")
        emit("")
        emit("| Version | Date | Author | Summary of Changes |")
        emit("|---|---|---|---|")
        emit("|  |  |  |  |")
        emit("")

    def emit_divider(section, role, slug):
        if subset:
            return
        if role == "procedure":
            divider_state["in_procedures"] = True
            l2 = l2_of.get(slug, "")
            if l2 and l2 != divider_state["l2"]:
                divider_state["l2"] = l2
                title = l2.replace("-", " ").title()
                emit(f"# {l2_ordinal.get(l2, '')}. {title}".replace("# . ", "# "))
                emit("")
        elif not divider_state["in_procedures"]:
            # Front matter (Process Overview, the reader-orientation views)
            # gets the same chapter treatment, or it floats parentless above
            # chapter 1 in the TOC and on the page.
            if not divider_state["front"]:
                divider_state["front"] = True
                emit("# Introduction")
                emit("")
        elif not divider_state["backmatter"]:
            divider_state["backmatter"] = True
            emit("# Reference & Appendices")
            emit("")

    for section in _sections(assembled):
        role = _attr(section, "role")
        slug = _attr(section, "slug") or ""
        if subset and not (role == "procedure" and slug in slugs):
            continue
        # gap-log: open gaps are working-draft scaffolding. screenshot-index:
        # the evidence lives inline at its steps in a final render (captured
        # images embed; the scorecard counts any still pending), so the
        # status appendix never ships to a client.
        if mode == "final" and _attr(section, "derived_kind") in (
                "gap-log", "screenshot-index"):
            continue
        heading = (_attr(section, "heading") or "").strip()
        raw_body = _attr(section, "body") or ""
        if _attr(section, "derived_kind") == "raci":
            # BEFORE token resolution, which replaces the [[slug]] sort keys.
            raw_body = _order_raci_rows(raw_body, numbers)
        if role == "procedure":
            submap = ids_by_slug.get(slug, {})
            if submap:
                raw_body = _rewrite_callout_ids(raw_body, submap)
            # The profile decides shape BEFORE final mode counts what it
            # strips, so a hidden section's gaps are never double-reported.
            raw_body = _apply_profile(raw_body, profile)
            # M16: the two body-view projections. Both are no-ops on a fragment
            # that uses none of the new grammar, and both run AFTER the profile
            # so a hidden section's callouts are already blank.
            raw_body = _hide_callout_details(raw_body)
            raw_body = _lead_conditions(raw_body)
        if mode == "final" and role == "procedure":
            raw_body = _final_transform(
                raw_body, slug, folder,
                disp_to_local_by_slug.get(slug, {}), stats)
        if mode == "final":
            # Citation scrub — EVERY section (derived registers and static
            # front/back matter cite sources too, not just procedures), after
            # every other body transform so it sees what the reader would.
            raw_body = _scrub_citations(raw_body, stats)
            # What the scrub cannot reach: an id woven into sentence meaning
            # ("... see GAP-37, which is unresolved") whose DEFINITION —
            # callout, gap-log row, sources.yaml entry — never reaches the
            # final reader. Reconcile cannot flag these (in the FRAGMENTS the
            # ids are still defined), and meaning-bearing prose cannot be
            # rewritten mechanically without risking the deliverable's
            # wording — so they are counted and enumerated for the human
            # running the export. Gap refs also clean themselves up through
            # the review round: closing a gap deletes its callout, at which
            # point the same prose mention DOES dangle at reconcile and the
            # drafter must remove it.
            refs = _GAP_ID_RE.findall(raw_body)
            if refs:
                key = slug or heading or "front-matter"
                stats["dangling_gap_refs"][key] = sorted(set(refs))
                stats["dangling_gap_ref_count"] += len(refs)
        if homeless_re is not None and role == "procedure":
            # Scanned AFTER every body transform, so it counts what the
            # reader of THIS mode actually sees (a final-mode strip can
            # remove a gap callout whose note cited a control).
            refs = homeless_re.findall(raw_body)
            if refs:
                stats["dangling_refs"][slug] = sorted(set(refs))
                stats["dangling_ref_count"] += len(refs)
        body = _clean_body(raw_body, labels)
        if mode == "final":
            body = body.strip("\n")
        derived_kind = _attr(section, "derived_kind")
        if mode == "final" and derived_kind:
            # Register grooming: the italic reviewer lead-in goes, and a
            # register with no rows at all is skipped whole rather than
            # shipping a heading over an empty shell.
            body = _strip_register_leadin(body).strip("\n")
            if _table_data_rows(body) == 0:
                stats["empty_registers_skipped"].append(heading)
                continue
        if mode == "final":
            body = _apply_lexicon(body, lexicon_normalizers, stats)
            _readiness_scan(body, slug or heading or "front-matter",
                            derived_kind, stats["readiness"])
        is_profile = heading.lower() in _PROFILE_HEADINGS

        # Lift the Document Profile onto the cover card (only when a cover is
        # being built); otherwise it renders inline like any other section.
        if do_cover and is_profile:
            profile_md = body
            continue

        first_backmatter = (divider_state["in_procedures"]
                            and not divider_state["backmatter"]
                            and role != "procedure")
        emit_divider(section, role, slug)
        if divider_state["backmatter"] and role != "procedure" \
                and not first_backmatter:
            break_headings.add(_heading_for(section, numbers))
        emit(f"## {_heading_for(section, numbers)}")
        emit("")
        if body.strip("\n"):
            fname = _attr(section, "file")
            trackable = role in ("procedure", "static")
            for k, ln in enumerate(body.split("\n")):
                emit(ln, {"file": fname, "slug": slug or None,
                          "l2": l2_of.get(slug, ""), "line": k}
                     if (provenance and trackable and fname) else None)
            emit("")

    # M30 — Shared Reference appendix: when this area's engagement carries
    # registers with CITABLE entries, they compile into the deliverable at
    # the very end of the body, so a "per the engagement approval matrix"
    # pointer finally resolves to something. CONTEXT entries never render —
    # they are align-only engagement intelligence, not deliverable content.
    # Kit (subset) renders skip it, same lean-kit rule as the L2 dividers;
    # both working and final modes carry it (citable = publishable).
    if not subset:
        parent = folder.resolve().parent
        if parent.name == "components":
            import registers as registers_mod
            blocks: list[tuple[str, list]] = []
            for _rp, rtitle, rentries in registers_mod.load_all(parent):
                cit = [e for e in (rentries or []) if e.cls == "citable"]
                if cit:
                    blocks.append((rtitle, cit))
            if blocks:
                if divider_state["in_procedures"] \
                        and not divider_state["backmatter"]:
                    emit("# Reference & Appendices")
                    emit("")
                else:
                    break_headings.add("Shared Reference — Engagement Registers")
                emit("## Shared Reference — Engagement Registers")
                emit("")
                for rtitle, cit in blocks:
                    emit(f"### {rtitle}")
                    emit("")
                    for e in cit:
                        # One paragraph per entry; the fragments' hard wraps
                        # rejoin so the converter sees a single run.
                        text = " ".join(ln.strip()
                                        for ln in e.text.split("\n")
                                        if ln.strip())
                        emit(f"**{e.id}** — {text}")
                        emit("")

    if prov is not None:
        prov.origins = origins

    body_md = "\n".join(body_lines)
    doc_id = uuid.uuid4().hex[:12]
    if mode == "final" and do_cover and profile_md.strip():
        # Cover readiness: a profile-card row shipping a blank or TBD value.
        for ln in profile_md.split("\n"):
            s = ln.strip()
            if not s.startswith("|") or _TABLE_SEP_RE.match(ln):
                continue
            cells = [c.strip() for c in s.strip("|").split("|")]
            if len(cells) >= 2 and cells[0].lower() != "field" \
                    and cells[1].upper() in ("", "TBD", "—", "-"):
                stats["readiness"]["cover_fields"].append(cells[0])
    cfgi.convert_assembled(
        body_md, out, title=title, subtitle=subtitle, profile_md=profile_md,
        include_toc=not subset, landscape=landscape, do_cover=do_cover,
        prov=prov, track_changes=track_changes, break_headings=break_headings,
    )
    stats["docx"] = str(out)

    if prov is not None and prov.entries:
        # Stamp the doc id into core properties so a RETURNED copy finds its
        # map no matter what the file got renamed to along the way.
        from docx import Document as _Doc
        d = _Doc(str(out))
        d.core_properties.category = f"cw-map:{doc_id}"
        d.save(str(out))
        basis = {}
        for comp in manifest.get("components", []):
            fp = folder / comp.get("file", "")
            if comp.get("role") in ("procedure", "static") and fp.is_file():
                basis[comp["file"]] = hashlib.sha1(fp.read_bytes()).hexdigest()
        map_path = _write_review_map(folder, out, doc_id, prov, basis)
        stats["doc_id"] = doc_id
        stats["map"] = str(map_path)

    # M7 signal file (full WORKING render only — M21: final mode is a
    # terminal export produced FROM an accepted state, not a transition into
    # a new one, so it must never touch the review signal — same reasoning
    # as the --slugs kit-subset skip below. Programmatic calls that pass
    # emit_signal=False are unaffected either way.)
    if emit_signal and not subset and mode == "working":
        import orchestrate
        orchestrate.emit_render(str(folder), str(out), awaiting_review=True)
    return stats


# --------------------------------------------------------------------------- #
# Input resolution
# --------------------------------------------------------------------------- #
def _resolve_input(area: str):
    """Return (kind, path). kind is 'folder' or 'file'."""
    p = Path(area)
    candidates = [p, _REPO_ROOT / "components" / area, Path("components") / area]
    for c in candidates:
        if c.is_dir() and (c / "manifest.json").is_file():
            return "folder", c
    if p.is_dir():
        # A directory without a manifest is not a valid area.
        raise SystemExit(f"error: no manifest.json in folder: {p}")
    if p.is_file():
        return "file", p
    raise SystemExit(f"error: input not found (no folder/manifest or .md): {area}")


def _infer_output(kind: str, path: Path, output_arg) -> Path:
    if output_arg:
        return Path(output_arg)
    if kind == "folder":
        name = _attr(doc_model.load_manifest(path), "area") or path.name
        return path / f"{name}_process-doc.docx"
    return cfgi.infer_output(path, None)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Render a CONSULT component folder (or single .md) to a CFGI-styled .docx")
    ap.add_argument("area", help="area folder (containing manifest.json) or a single .md file")
    ap.add_argument("-o", "--output")
    ap.add_argument("--mode", choices=("working", "final"), default="working",
                    help="working: everything visible + provenance; final: client-facing (gaps stripped, screenshots embedded)")
    ap.add_argument("--slugs", default=None,
                    help="comma-separated procedure slugs — subset (kit) render")
    ap.add_argument("--track-changes", action="store_true",
                    help="emit the docx with tracked changes on by default")
    ap.add_argument("--strict", action="store_true",
                    help="final mode: exit nonzero when the readiness "
                         "scorecard is not clean (the render still completes)")
    ap.add_argument("--include-toc", action="store_true",
                    help="single-file input only: insert a Table of Contents "
                         "(folder renders ALWAYS carry one; kit/subset "
                         "renders never do)")
    ap.add_argument("--landscape", action="store_true", help="Use landscape orientation")
    ap.add_argument("--no-cover", action="store_true", help="Skip the generated cover page")
    a = ap.parse_args(argv)

    kind, path = _resolve_input(a.area)
    out = _infer_output(kind, path, a.output)
    do_cover = not a.no_cover
    if kind == "folder":
        slugs = [s.strip() for s in a.slugs.split(",") if s.strip()] if a.slugs else None
        stats = render_folder(
            path, out, landscape=a.landscape,
            do_cover=do_cover, mode=a.mode, slugs=slugs,
            track_changes=a.track_changes)
        print("Wrote " + str(out))
        print(stats["profile"])
        if a.mode == "final":
            print(f"final mode: stripped {stats['gaps_stripped']} open gap callout(s) "
                  f"+ {stats['gap_tags_stripped']} inline tag(s); "
                  f"scrubbed {stats['citations_scrubbed']} SRC/GAP citation(s); "
                  f"{stats['screens_embedded']} screenshot(s) embedded, "
                  f"{stats['screens_placeholder']} placeholder(s) remain")
            if stats["lexicon_normalized"]:
                print(f"lexicon: normalized {stats['lexicon_normalized']} "
                      f"spelling variant(s) to canonical form")
            if stats["empty_registers_skipped"]:
                print("skipped empty register(s): "
                      + ", ".join(stats["empty_registers_skipped"]))
            # Readiness scorecard — counts reported, never refused.
            r = stats["readiness"]
            labels = [("cover_fields", "blank cover field(s)"),
                      ("placeholders", "placeholder string(s)"),
                      ("blank_cells", "blank required register cell(s)"),
                      ("double_spaces", "doubled space(s) in prose "
                       "(possible dropped word)"),
                      ("double_punct", "doubled punctuation mark(s)")]
            dirty = sum(len(r[k]) for k, _ in labels) \
                + stats["screens_placeholder"]
            if dirty:
                print(f"READINESS: {dirty} item(s) need attention before "
                      f"client delivery:")
                if stats["screens_placeholder"]:
                    print(f"  - {stats['screens_placeholder']} screenshot(s) "
                          f"still pending capture (placeholder rendered)")
                for key, label in labels:
                    for item in r[key]:
                        print(f"  - {label}: {item}")
            else:
                print("READINESS: clean — no placeholders, blank required "
                      "cells, or punctuation defects detected")
            if stats["dangling_gap_ref_count"]:
                print(f"WARNING: {stats['dangling_gap_ref_count']} SRC/GAP "
                      f"reference(s) survive in ordinary prose — they are "
                      f"woven into sentence meaning, so the citation scrub "
                      f"could not remove them mechanically, and what they "
                      f"point at is not in this final render:")
                for pslug, ids in sorted(stats["dangling_gap_refs"].items()):
                    print(f"  {pslug}: " + ", ".join(ids))
                print("  close gaps through the review round (once a gap "
                      "closes, its leftover prose reference fails reconcile "
                      "until the drafter removes it), or hand-edit the prose "
                      "before shipping this export")
        # Profile-shaped, so it applies in EVERY mode — a working draft
        # without a home for these kinds has the same holes its final would.
        if stats["dangling_ref_count"]:
            print(f"WARNING: {stats['dangling_ref_count']} dangling callout "
                  f"reference(s) survive in prose — this profile gives their "
                  f"kind no home (no section in the body, no register in the "
                  f"manifest):")
            for pslug, ids in sorted(stats["dangling_refs"].items()):
                print(f"  {pslug}: " + ", ".join(ids))
            print("  keep the section, or hide it with `body_omit` plus its "
                  "register (scaffold.py --sync-profile adds a register to a "
                  "scaffolded area), or reword the prose before shipping")
        if a.strict and a.mode == "final" and (
                dirty or stats["dangling_gap_ref_count"]
                or stats["dangling_ref_count"]):
            print("--strict: readiness scorecard is not clean, exiting 1 "
                  "(the render above completed and is on disk)")
            return 1
    else:
        if a.mode != "working" or a.slugs or a.track_changes:
            raise SystemExit("error: --mode/--slugs/--track-changes require an area folder")
        cfgi.convert(path, out, a.include_toc, a.landscape, do_cover)
        print("Wrote " + str(out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
