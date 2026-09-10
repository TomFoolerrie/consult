#!/usr/bin/env python3
"""render_worker — the one Python seam: a bounded docx formatter that never thinks.

Reads ONE versioned JSON job (v1) on stdin, writes a .docx at job["out"],
prints {"path", "sections", "warnings"} on stdout. Every content decision
(what the sections say, their order, the title, whether it is a draft) was
made in TypeScript before the job arrived; this file only formats.

Job v1:
  { "version": 1, "title": str, "draft": bool,
    "skin": { "format": "docx", "requires": [str] },
    "sections": [ { "id": str, "title": str, "body": str } ],
    "out": str }

Body markup understood (what the view builders emit):
  "- item"      -> a List Bullet paragraph
  "### heading" -> a Heading 2
  blank line    -> paragraph break
  anything else -> a Normal paragraph (consecutive lines join one paragraph)

Exits: 0 on success · 2 on a job this worker cannot accept (JSON error on
stdout) · 1 on any other failure (message on stderr). No reading of the
engagement folder, no network, no judgment.
"""
import json
import sys

VERSION = 1
DRAFT_LINE = "DRAFT — not for distribution"


def refuse(msg: str) -> None:  # exit 2: a job we decline, said in JSON so the caller can name it
    print(json.dumps({"error": msg}))
    sys.exit(2)


def add_body(doc, body: str, warnings: list) -> None:
    """render the light markup into paragraphs; unknown constructs degrade to Normal text with a warning"""
    para_lines: list = []

    def flush() -> None:
        if para_lines:
            doc.add_paragraph(" ".join(para_lines), style="Normal")
            para_lines.clear()

    for raw in (body or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            flush()
        elif line.startswith("### "):
            flush()
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            flush()
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("#"):
            flush()
            warnings.append(f"unsupported heading depth rendered as text: {line[:40]}")
            doc.add_paragraph(line.lstrip("# ").strip(), style="Normal")
        else:
            para_lines.append(line.strip())
    flush()


def render(job: dict) -> dict:
    from docx import Document
    from docx.enum.text import WD_BREAK

    warnings: list = []
    doc = Document()
    title = str(job.get("title") or "")
    draft = bool(job.get("draft", False))
    requires = list((job.get("skin") or {}).get("requires") or [])
    sections = list(job.get("sections") or [])

    if draft:
        for section in doc.sections:
            section.header.paragraphs[0].text = "DRAFT"
    doc.core_properties.title = title

    if "title-page" in requires:
        doc.add_paragraph(title, style="Title")
        if draft:
            doc.add_paragraph(DRAFT_LINE)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for s in sections:
        doc.add_heading(str(s.get("title") or s.get("id") or ""), level=1)
        add_body(doc, str(s.get("body") or ""), warnings)
        if not (s.get("body") or "").strip():
            warnings.append(f"section {s.get('id')} has an empty body")

    for cap in requires:
        if cap != "title-page":
            warnings.append(f"skin capability not understood by this worker: {cap}")

    doc.save(job["out"])
    return {"path": job["out"], "sections": len(sections), "warnings": warnings}


def main() -> int:
    try:
        job = json.load(sys.stdin)
    except json.JSONDecodeError as e:
        refuse(f"job is not JSON: {e}")
    if not isinstance(job, dict) or job.get("version") != VERSION:
        refuse(f"job version {job.get('version') if isinstance(job, dict) else '?'} not supported — this worker formats v{VERSION}")
    if (job.get("skin") or {}).get("format", "docx") != "docx":
        refuse(f"skin format {job['skin'].get('format')} not supported — this worker emits docx")
    if not job.get("out"):
        refuse("job names no output path")
    try:
        print(json.dumps(render(job)))
        return 0
    except Exception as e:  # noqa: BLE001 — every failure surfaces as exit 1 + one message
        print(f"{type(e).__name__}: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
